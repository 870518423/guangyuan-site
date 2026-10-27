# -*- coding: utf-8 -*-
"""
一键生成「奥系人间体大全」查询页
================================================
用法：
    python build_guides.py                     # 用默认表格路径
    python build_guides.py "D:\\某处\\人间体大全.xlsx"

流程：
    读 Excel（4 个阵营子表） → 解析成 JSON → 填入 guide_template.html → 输出 guides.html

表格列约定（4 张子表一致，表头第 1 行，数据从第 2 行开始）：
    A 人间体名称 | B 本命 | C 获得方式 | D 前置要求 | E~I 特训人间体（最多 5 个）

改完表格后重跑本脚本即可，不需要手工改网页。
"""
import json
import os
import sys
import datetime

try:
    import openpyxl
except ImportError:
    sys.exit("缺少依赖：请先安装 openpyxl（pip install openpyxl）")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_SRC = r"C:\Users\Admin\Desktop\奥系\人间体大全（含特训及条件）.xlsx"
TEMPLATE = os.path.join(HERE, "guide_template.html")
OUT = os.path.join(HERE, "guides.html")

# ---- 人间体图片（2026-10-17 新增）----
# 图片由 tools/build_human_imgs.py 压好落在 images/humans/<id>/，清单同名 manifest.json
IMG_MANIFEST = os.path.join(HERE, "images", "humans", "manifest.json")
# Excel 用简体正字、图片目录按游戏内写法 → 少量别名（其余已把 Excel 改对）
IMG_ALIAS = {"姬矢准": "姫矢准"}

# 性别判定表（2026-10-17）：A 序号 / B 角色名 / C 形态数 / D 性别 / E 依据
SEX_SRC = r"C://Users//Admin//Desktop//奥系//奥系人间体//人间体性别判定表.xlsx"


def load_sex():
    """读性别判定表 → {角色名: 男/女/不明}；文件缺失就返回空（构建照常，只是不显示性别）。"""
    if not os.path.isfile(SEX_SRC):
        print("提示：没找到性别判定表，本次生成不带性别。")
        return {}
    wb = openpyxl.load_workbook(SEX_SRC, data_only=True)
    ws = wb.worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=3, values_only=True):   # 第 1 行说明、第 2 行表头
        if not row or len(row) < 4:
            continue
        name, sex = row[1], row[3]
        if name and sex:
            out[str(name).strip()] = str(sex).strip()
    return out


def load_img_manifest():
    """读图片清单；没生成过就返回空（构建照常，只是不带图）。"""
    if not os.path.isfile(IMG_MANIFEST):
        print("提示：还没有图片清单（先跑 tools/build_human_imgs.py），本次生成不带图。")
        return {}
    with open(IMG_MANIFEST, "r", encoding="utf-8") as f:
        return json.load(f).get("chars", {})

# 阵营 → 图标 / 主题色（改图标只需换这里的文件名）
FACTIONS = [
    {"name": "奥特之星", "icon": "assets/factions/star.png",       "color": "#fbbf24"},
    {"name": "多宇宙",   "icon": "assets/factions/multiverse.png", "color": "#38bdf8"},
    {"name": "怪兽",     "icon": "assets/factions/monster.png",    "color": "#34d399"},
    {"name": "宇宙人",   "icon": "assets/factions/alien.png",      "color": "#a78bfa"},
]

# 多值分隔符：换行 + 全角逗号 + 顿号
SEPS = "\n，、"

COL_NAME, COL_FATE, COL_HOW, COL_PREQ = 0, 1, 2, 3
TRAIT_RANGE = (4, 9)  # E~I

# 礼物按品质从高到低排列；表里可以写完整名称，也可以写前两个字缩写。
GIFTS_BY_QUALITY = (
    ("格斗手套", "恶魔之花", "宇宙地图", "白色围巾"),
    ("棒球手套", "能量指环", "咖喱饭", "光子计算机"),
    ("战斗日记", "飞行执照", "巴顿羽毛"),
)
GIFT_ORDER = {name: i for i, name in enumerate(
    gift for quality in GIFTS_BY_QUALITY for gift in quality
)}
GIFT_NAMES_BY_LENGTH = sorted(GIFT_ORDER, key=len, reverse=True)
GIFT_ABBREVIATIONS = {name[:2]: name for name in GIFT_ORDER}
GIFT_ABBREVIATIONS.update({"围巾": "白色围巾", "羽毛": "巴顿羽毛"})

# 「本命 / 前置要求」里同名但不是人间体、而是「角色卡」的特例
# ------------------------------------------------------------------
# B 列（本命）和 D 列（前置要求）里绝大多数写的是角色卡名（赛文、盖亚、
# 雷欧、托雷基亚…），本来就不跟人间体重名、不会跳转；只有极少数会重名。
# 凡重名但其实指「角色卡」的，在下面按 (人间体名, 该项原文) 登记一行，
# 生成时会加「（角色卡）」后缀 → 名字不再等于任何人间体 → 自动失去跳转。
#
# 注意：E~I 列（特训所需）不在此列，特训那一栏本来就是指人间体。
#
# 新增例子：("某个间体", "某项本命或前置")
CARD_ITEMS = {
    ("斯菲亚虚体", "斯菲亚"),     # 前置：角色卡「斯菲亚」≠ 人间体「斯菲亚」
    ("雷布朗多遗族", "贝利亚"),   # 本命 + 前置都是角色卡「贝利亚」≠ 人间体「贝利亚」
    ("雾崎", "托雷基亚"),         # 本命：角色卡「托雷基亚」≠ 人间体「托雷基亚」
    ("卡露蜜拉", "卡露蜜拉"),     # 本命：角色卡「卡露蜜拉」≠ 人间体（自身）
    ("苏伊德", "苏伊德"),         # 本命：角色卡「苏伊德」≠ 人间体（自身）
    ("天照女王", "天照女王"),     # 本命：角色卡「天照女王」≠ 人间体（自身）
}
CARD_SUFFIX = "（角色卡）"
CARD_FIELDS = ("fate", "prereq")   # 生效的列：本命 + 前置要求


def split_multi(value):
    """把一格里的多值拆成列表，去空、去重、保留顺序。"""
    if value is None:
        return []
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    parts = []
    for chunk in text.split("\n"):
        for piece in chunk.replace("，", "\u0001").replace("、", "\u0001").split("\u0001"):
            piece = piece.strip().strip("，、,;；").strip()
            if piece and piece not in parts:
                parts.append(piece)
    return parts


def normalize_gifts(value):
    """把礼物缩写还原为完整名称，并按紫、蓝、绿品质排序。"""
    if value is None:
        return []
    text = str(value)
    for ch in "\r\n，、,;； \t":
        text = text.replace(ch, "")
    if not text:
        return []

    gifts = []
    pos = 0
    while pos < len(text):
        name = next((g for g in GIFT_NAMES_BY_LENGTH if text.startswith(g, pos)), None)
        if name:
            gifts.append(name)
            pos += len(name)
            continue
        abbr = text[pos:pos + 2]
        name = GIFT_ABBREVIATIONS.get(abbr)
        if name:
            gifts.append(name)
            pos += 2
            continue
        raise ValueError(f"无法识别礼物「{text[pos:]}」，原内容：{value}")

    return sorted(dict.fromkeys(gifts), key=GIFT_ORDER.get)


def load_entries(src):
    wb = openpyxl.load_workbook(src, data_only=True)
    known = {f["name"] for f in FACTIONS}
    pics = load_img_manifest()
    sexes = load_sex()
    entries = []
    per_faction = {}

    for ws in wb.worksheets:
        title = ws.title.strip()
        faction = title if title in known else title
        header = {}
        for cell in ws[1]:
            value = str(cell.value or "").strip()
            if value:
                header[value] = cell.column - 1
        for required in ("喜欢的礼物", "讨厌的礼物"):
            if required not in header:
                sys.exit(f"[{title}] 缺少表头：{required}")
        like_col = header["喜欢的礼物"]
        dislike_col = header["讨厌的礼物"]
        got = 0
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                min_col=1, max_col=ws.max_column,
                                values_only=True):
            name = row[COL_NAME]
            if name is None or not str(name).strip():
                continue
            name = str(name).strip()
            traits = []
            for cell in row[TRAIT_RANGE[0]:TRAIT_RANGE[1]]:
                for t in split_multi(cell):
                    if t not in traits:
                        traits.append(t)
            try:
                like_gifts = normalize_gifts(row[like_col])
                dislike_gifts = normalize_gifts(row[dislike_col])
            except ValueError as exc:
                sys.exit(f"[{title}] {name}：{exc}")

            # 本命 / 前置：命中 CARD_ITEMS 的加「（角色卡）」后缀，避免误链到人间体
            def mark(values):
                out = []
                for v in values:
                    if (name, v) in CARD_ITEMS:
                        v = v + CARD_SUFFIX
                    if v not in out:
                        out.append(v)
                return out

            fate = split_multi(row[COL_FATE])
            prereq = split_multi(row[COL_PREQ])
            pic = pics.get(IMG_ALIAS.get(name, name))
            sex = sexes.get(name) or sexes.get(IMG_ALIAS.get(name, name), "")
            if "fate" in CARD_FIELDS:
                fate = mark(fate)
            if "prereq" in CARD_FIELDS:
                prereq = mark(prereq)

            entries.append({
                "name": name,
                "faction": faction,
                "fate": fate,
                "how": split_multi(row[COL_HOW]),
                "prereq": prereq,
                "traits": traits,
                "likeGifts": like_gifts,
                "dislikeGifts": dislike_gifts,
                "img": pic["id"] if pic else "",
                "forms": pic["forms"] if pic else [],
                "av": 1 if (pic and pic.get("avatar")) else 0,
                "sex": sex,
            })
            got += 1
        per_faction[faction] = got
        print(f"  [{title}] {got} 条")

    return entries, per_faction

def find_source(explicit):
    """定位表格：优先用命令行传入的路径；否则用默认路径；再否则自动在桌面里找。"""
    if explicit:
        if os.path.isfile(explicit):
            return explicit
        sys.exit(f"找不到表格文件：{explicit}")
    if os.path.isfile(DEFAULT_SRC):
        return DEFAULT_SRC
    # 默认路径失效时（比如文件被改名），在桌面里自动找最新的「人间体*.xlsx」
    cands = []
    for folder in (os.path.join(os.path.expanduser("~"), "Desktop", "奥系"),
                   os.path.join(os.path.expanduser("~"), "Desktop"),
                   os.path.join(os.path.expanduser("~"), "Desktop", "奥系道具图片")):
        if not os.path.isdir(folder):
            continue
        for fn in os.listdir(folder):
            low = fn.lower()
            if low.endswith(".xlsx") and "人间体" in fn and not fn.startswith("~$"):
                cands.append(os.path.join(folder, fn))
    if cands:
        cands.sort(key=os.path.getmtime, reverse=True)
        picked = cands[0]
        print(f"注意：默认表格不存在，已自动改用：{picked}")
        if len(cands) > 1:
            print("      （同目录下还有其它候选，如需指定请把表格拖到「生成攻略页.bat」上）")
        return picked
    sys.exit(
        "找不到表格文件。\n"
        f"  默认路径：{DEFAULT_SRC}\n"
        "  解决办法：把 Excel 表格直接拖到「生成攻略页.bat」上运行，"
        "或在命令行里把路径作为参数传进来。"
    )


def main():
    src = find_source(sys.argv[1] if len(sys.argv) > 1 else None)
    if not os.path.isfile(TEMPLATE):
        sys.exit(f"找不到模板文件：{TEMPLATE}")

    print("读取表格：", src)
    entries, per_faction = load_entries(src)
    if not entries:
        sys.exit("没有解析到任何条目，请检查表格结构。")

    data = {
        "generated": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source": os.path.basename(src),
        "factions": FACTIONS,
        "entries": entries,
    }
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))

    with open(TEMPLATE, "r", encoding="utf-8") as fh:
        tpl = fh.read()
    if "/*__GUIDE_DATA__*/" not in tpl:
        sys.exit("模板里找不到数据占位符 /*__GUIDE_DATA__*/")
    html = tpl.replace("/*__GUIDE_DATA__*/ null", payload)

    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write(html)

    size = os.path.getsize(OUT)
    print("\n生成完成：", OUT)
    print(f"  条目总数：{len(entries)}")
    for f in FACTIONS:
        print(f"    {f['name']}: {per_faction.get(f['name'], 0)}")
    print(f"  输出体积：{size:,} 字节（约 {size/1024:.1f} KB）")
    no_sex = [e["name"] for e in entries if not e.get("sex")]
    if no_sex:
        print(f"  提示：{len(no_sex)} 条没有性别 → {', '.join(no_sex[:6])}")
    no_trait = [e["name"] for e in entries if not e["traits"]]
    if no_trait:
        print(f"  提示：{len(no_trait)} 条没有任何特训人间体 → {', '.join(no_trait[:8])}")


if __name__ == "__main__":
    main()
